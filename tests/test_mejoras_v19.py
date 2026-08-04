# -*- coding: utf-8 -*-
"""Tests de las mejoras v19 — contra el CÓDIGO REAL del disco.

Cubren:
  M1  websearch: caché con TTL + extract_text (article/h/p/li) + research unificado
  M2  briefing: briefing_due (control horario), secciones tablero/hermes, coach integrado
  M3  vigilancias: precio/huella/diff, alta numerada, detección de cambios y avisos
  M5  informes: md_to_docx real (python-docx) + histórico y reapertura
  M10 backup: zip de data/ + rotación + archivado verificado de .bak

Ejecutar:  python tests/test_mejoras_v19.py    (desde la carpeta nexus)
"""
import asyncio
import datetime as dt
import importlib.util
import os
import re
import sys
import tempfile
import time
import types
import zipfile
from pathlib import Path

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_fail = []
_pass = 0


def check(cond, msg):
    global _pass
    if cond:
        _pass += 1
    else:
        _fail.append(msg)
        print("  FALLO:", msg)


def _stub(name):
    m = types.ModuleType(name)
    m.__getattr__ = lambda _n: types.SimpleNamespace()
    sys.modules[name] = m
    return m


def load_skill_module(folder):
    if ROOT not in sys.path:
        sys.path.insert(0, ROOT)
    path = os.path.join(ROOT, "skills", folder, "skill.py")
    for _ in range(15):
        spec = importlib.util.spec_from_file_location(f"_v19_{folder}", path)
        mod = importlib.util.module_from_spec(spec)
        try:
            spec.loader.exec_module(mod)
            return mod
        except ModuleNotFoundError as e:
            _stub(e.name)
    raise RuntimeError(folder)


def _tmp_datadir():
    """Redirige backend.core.comun.config.DATA_DIR a una carpeta temporal."""
    if ROOT not in sys.path:
        sys.path.insert(0, ROOT)
    import backend.core.comun.config as cfg
    tmp = Path(tempfile.mkdtemp())
    old = cfg.DATA_DIR
    cfg.DATA_DIR = tmp
    return cfg, tmp, old


# ══════════════ M1: websearch — caché TTL + extracción ══════════════

def test_websearch_cache():
    cfg, tmp, old = _tmp_datadir()
    try:
        import backend.core.infraestructura.websearch as ws
        ws.cache_put("q::prueba::6", [{"title": "T", "snippet": "S", "url": "u"}])
        hit = ws.cache_get("q::prueba::6", ttl=60)
        check(hit and hit[0]["title"] == "T", "caché: put/get devuelve lo guardado")
        # TTL caducado → None (reescribimos el ts a mano)
        d = ws._cache_load()
        d["q::prueba::6"]["ts"] = time.time() - 999
        ws._cache_save(d)
        check(ws.cache_get("q::prueba::6", ttl=60) is None, "caché: TTL caducado → None")
        check(ws.cache_get("q::no-existe", ttl=60) is None, "caché: miss → None")
        # search() sirve desde caché SIN red (si tocara la red con los stubs, devolvería [])
        ws.cache_put("q::quién ganó ayer::6", [{"title": "Cacheado", "snippet": "", "url": ""}])
        res = asyncio.run(ws.search("quién ganó ayer", 6))
        check(res and res[0]["title"] == "Cacheado", "search(): sirve desde caché sin tocar la red")
        # poda: nunca más de _CACHE_MAX entradas
        for i in range(ws._CACHE_MAX + 20):
            ws.cache_put(f"q::masivo{i}", [i])
        check(len(ws._cache_load()) <= ws._CACHE_MAX, "caché: poda a _CACHE_MAX entradas")
    finally:
        cfg.DATA_DIR = old


def test_websearch_extract():
    import backend.core.infraestructura.websearch as ws
    html = """
    <html><head><script>var x=1;</script><style>.a{}</style></head><body>
    <nav><li>Menú uno de navegación que no interesa</li></nav>
    <article>
      <h1>Titular principal del artículo</h1>
      <p>Primer párrafo con contenido de verdad y suficiente longitud para contar.</p>
      <p>corto</p>
      <li>Elemento de lista informativo con datos interesantes del tema tratado</li>
    </article>
    <footer><p>Pie de página legal que no debería salir porque está en footer.</p></footer>
    </body></html>"""
    text = ws.extract_text(html)
    check("Titular principal" in text, "extract: capta h1 dentro de article")
    check("Primer párrafo" in text, "extract: capta párrafos con sustancia")
    check("Pie de página legal" not in text, "extract: descarta footer")
    check("var x=1" not in text, "extract: descarta scripts")
    check("Menú uno" not in text, "extract: descarta nav")
    check("corto" not in text.replace("Titular", ""), "extract: descarta párrafos triviales")
    # sin estructura → texto plano
    check("hola mundo" in ws.extract_text("<div>hola mundo</div>" + "x" * 10), "extract: fallback plano")


def test_research_usa_motor_central():
    src = open(os.path.join(ROOT, "skills", "research", "skill.py"), encoding="utf-8").read()
    check("websearch.search" in src and "websearch.fetch_page" in src,
          "research: delega en backend.core.infraestructura.websearch")
    check("html.duckduckgo.com" not in src, "research: sin scraping duplicado de DDG")


# ══════════════ M2: briefing ══════════════

def test_briefing_due():
    import backend.core.dominio.briefing as bf
    now = dt.datetime(2026, 7, 24, 9, 0)
    check(bf.briefing_due(now, True, "08:30", "") is True, "due: pasada la hora y sin mandar → sí")
    check(bf.briefing_due(now, True, "08:30", "2026-07-24") is False, "due: ya mandado hoy → no")
    check(bf.briefing_due(now, True, "09:30", "") is False, "due: aún no es la hora → no")
    check(bf.briefing_due(now, False, "08:30", "") is False, "due: desactivado → no")
    check(bf.briefing_due(now, True, "basura", "") is False, "due: hora inválida → no revienta")


def test_briefing_secciones():
    import json
    import backend.core.dominio.briefing as bf
    tmp = Path(tempfile.mkdtemp())
    old = bf.DATA_DIR
    bf.DATA_DIR = tmp
    try:
        # hermes: activos y recientes con número
        (tmp / "hermes_jobs.json").write_text(json.dumps([
            {"num": 1, "estado": "hecho", "t1": time.time() - 100, "orden": "x"},
            {"num": 2, "estado": "trabajando", "t0": time.time(), "orden": "y"},
        ]), encoding="utf-8")
        s = bf._sec_hermes()
        check("#2" in s and "#1" in s, "briefing hermes: menciona encargos por número")
        (tmp / "hermes_jobs.json").write_text("[]", encoding="utf-8")
        check(bf._sec_hermes() == "", "briefing hermes: sin encargos → sección vacía")
    finally:
        bf.DATA_DIR = old
    # el scheduler llama a maybe_send y coach usa build_briefing
    sch = open(os.path.join(ROOT, "backend", "core", "scheduler.py"), encoding="utf-8").read()
    check("briefing.maybe_send" in sch, "scheduler: engancha el briefing")
    coach = open(os.path.join(ROOT, "skills", "coach", "skill.py"), encoding="utf-8").read()
    check("build_briefing" in coach, "coach: «qué me toca hoy» usa el parte completo")


# ══════════════ M3: vigilancias ══════════════

def test_vigilancias_detectores():
    v = load_skill_module("vigilancias")
    check(v.extract_price("cuesta 1.234,56 € con envío") == 1234.56, "precio: 1.234,56 €")
    check(v.extract_price("oferta €99 hoy") == 99.0, "precio: €99")
    check(v.extract_price("son 120 euros al mes") == 120.0, "precio: 120 euros")
    check(v.extract_price("price: $45.50") == 45.5, "precio: $45.50")
    check(v.extract_price("no hay nada aquí") is None, "precio: sin precio → None")
    d1, d2 = v.text_digest("hola  mundo"), v.text_digest("HOLA MUNDO")
    check(d1 == d2, "digest: normaliza espacios y mayúsculas")
    check(v.text_digest("otra cosa") != d1, "digest: contenido distinto → huella distinta")
    fresh = v.diff_lines("línea antigua uno\nlínea antigua dos",
                         "línea antigua uno\nesta línea es completamente nueva y larga")
    check(fresh == ["esta línea es completamente nueva y larga"], "diff: detecta solo lo nuevo")


def test_vigilancias_ciclo():
    v = load_skill_module("vigilancias")
    cfg, tmp, old = _tmp_datadir()
    avisos = []

    async def fake_notify(msg):
        avisos.append(msg)
    v._notify = fake_notify
    # stub del motor web: contenido controlado que CAMBIA entre pasadas
    state = {"page": "contenido inicial estable con su precio de 100 euros aquí"}

    fake_ws = types.ModuleType("backend.core.infraestructura.websearch")

    async def fake_fetch(url, max_chars=6000):
        return state["page"]

    async def fake_search(q, n=5):
        return [{"title": t} for t in state.get("titulares", [])]
    fake_ws.fetch_page = fake_fetch
    fake_ws.search = fake_search
    old_ws = sys.modules.get("backend.core.infraestructura.websearch")
    sys.modules["backend.core.infraestructura.websearch"] = fake_ws
    # El doble hay que colgarlo del PAQUETE que hace el import, y websearch bajo
    # a backend/core/infraestructura/ con la Fase 3. Colgarlo del paquete viejo
    # dejaba el doble sin efecto y la prueba veia el motor de verdad.
    import backend.core.infraestructura as _bc
    _old_attr = getattr(_bc, "websearch", None)
    _bc.websearch = fake_ws
    try:
        w1 = v._add("web", "https://ejemplo.com")
        w2 = v._add("precio", "https://tienda.com/p")
        w3 = v._add("noticias", "calcetines técnicos")
        check((w1["num"], w2["num"], w3["num"]) == (1, 2, 3), "vigilancias numeradas 1,2,3")
        state["titulares"] = ["titular viejo A", "titular viejo B"]
        asyncio.run(v.check_watchers(force=True))          # 1ª pasada: siembra
        check(avisos == [], "1ª pasada siembra sin avisar")
        # 2ª pasada: cambia la web, baja el precio, sale titular nuevo
        state["page"] = "contenido totalmente distinto ahora con el precio rebajado a 80 euros ya"
        state["titulares"] = ["titular viejo A", "titular NUEVO que acaba de salir"]
        asyncio.run(v.check_watchers(force=True))
        joined = " || ".join(avisos)
        check("#1" in joined and "cambiado" in joined, "aviso de cambio web con número")
        check("BAJA" in joined and "80.00" in joined, "aviso de bajada de precio 100→80")
        check("titular NUEVO" in joined, "aviso de titular nuevo")
        # borrar por número
        rx = re.compile(v.SKILL["patterns"]["remove"], re.IGNORECASE)
        m = rx.search("borra la vigilancia 2")
        r = asyncio.run(v.handle("remove", "borra la vigilancia 2", m, {}))
        check("#2" in r["reply"] and "eliminada" in r["reply"], "borrado por número")
        check(len(v._load()) == 2, "quedan 2 vigilancias tras borrar")
    finally:
        if old_ws is not None:
            sys.modules["backend.core.infraestructura.websearch"] = old_ws
        if _old_attr is not None:
            _bc.websearch = _old_attr
        cfg.DATA_DIR = old


# ══════════════ M5: informes docx + histórico ══════════════

def test_informes_docx_e_historico():
    r = load_skill_module("research")
    tmp = Path(tempfile.mkdtemp())
    md = ("# Informe de prueba\n\n_meta_\n\n## Hallazgos\n\n"
          "- Primer hallazgo importante\n- Segundo hallazgo\n\nPárrafo normal de cierre.")
    out = tmp / "informe.docx"
    ok = r.md_to_docx(md, out)
    check(ok and out.exists() and out.stat().st_size > 1000, "md_to_docx crea un .docx real")
    if ok:
        import docx
        d = docx.Document(str(out))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        bullets = [p.text for p in d.paragraphs if p.style.name == "List Bullet"]
        check("Informe de prueba" in heads and "Hallazgos" in heads, "docx: encabezados md → Word")
        check("Primer hallazgo importante" in bullets, "docx: viñetas md → List Bullet")
    # histórico con carpeta temporal
    old_dir = r.REPORTS_DIR
    r.REPORTS_DIR = tmp
    try:
        (tmp / "proveedores-china-2026-07-20.md").write_text("# a", encoding="utf-8")
        (tmp / "mercado-calcetines-2026-07-24.md").write_text("# b", encoding="utf-8")
        res = asyncio.run(r.handle("history", "mis informes", None, {}))
        check("proveedores-china" in res["reply"] and "mercado-calcetines" in res["reply"],
              "histórico lista los informes")
        rx = re.compile(r.SKILL["patterns"]["open_report"], re.IGNORECASE)
        m = rx.search("abre el informe de calcetines")
        import webbrowser
        old_open = webbrowser.open
        webbrowser.open = lambda *_a, **_k: True
        try:
            res2 = asyncio.run(r.handle("open_report", "abre el informe de calcetines", m, {}))
        finally:
            webbrowser.open = old_open
        check("mercado-calcetines" in res2["reply"], "reabre el informe por nombre parcial")
        m3 = rx.search("abre el informe de marcianos")
        res3 = asyncio.run(r.handle("open_report", "abre el informe de marcianos", m3, {}))
        check("No encuentro" in res3["reply"], "informe inexistente → aviso claro")
    finally:
        r.REPORTS_DIR = old_dir


# ══════════════ M10: backup ══════════════

def test_backup():
    b = load_skill_module("backup")
    cfg, tmp, old = _tmp_datadir()
    try:
        # data/ con contenido y cosas a excluir
        (tmp / "memory").mkdir(parents=True)
        (tmp / "memory" / "nota.md").write_text("hola", encoding="utf-8")
        (tmp / "board.json").write_text("{}", encoding="utf-8")
        (tmp / "chrome_nexus").mkdir()
        (tmp / "chrome_nexus" / "perfil.bin").write_text("x", encoding="utf-8")
        out, n = b.make_backup()
        check(out.exists() and n == 2, f"backup: zip con 2 archivos (excluye perfil) — n={n}")
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
        check("memory/nota.md" in names and "board.json" in names, "backup: contenido correcto")
        check(not any("chrome_nexus" in x for x in names), "backup: excluye el perfil de Chrome")
        # rotación
        dest = tmp / "backups"
        for i in range(10):
            (dest / f"nexus-data-2026010{i}.zip").write_bytes(b"PK\x05\x06" + b"\x00" * 18)
        b.rotate_backups(keep=7)
        check(len(list(dest.glob("nexus-data-*.zip"))) == 7, "backup: rotación deja 7")
        # archivado de .bak con verificación (árbol de proyecto temporal)
        proj = Path(tempfile.mkdtemp())
        (proj / "data").mkdir()
        (proj / "skills" / "x").mkdir(parents=True)
        (proj / "skills" / "x" / "skill.py.bak_v9").write_text("viejo1", encoding="utf-8")
        (proj / "app.py.bak_v2").write_text("viejo2", encoding="utf-8")
        cfg.DATA_DIR = proj / "data"
        n2, out2 = b.archive_baks(proj)
        check(n2 == 2 and out2 and out2.exists(), f"baks: 2 archivados — n={n2}")
        check(not (proj / "app.py.bak_v2").exists(), "baks: originales borrados tras verificar")
        with zipfile.ZipFile(out2) as z:
            check(z.testzip() is None and len(z.namelist()) == 2, "baks: zip íntegro con los 2")
    finally:
        cfg.DATA_DIR = old
    sch = open(os.path.join(ROOT, "backend", "core", "scheduler.py"), encoding="utf-8").read()
    check("auto_backup" in sch and "check_watchers" in sch, "scheduler: engancha backup y vigilancias")


# ══════════════ Enrutado global de lo nuevo + anti-robo ══════════════

def test_routing_v19():
    sys.path.insert(0, os.path.join(ROOT, "tests"))
    from test_renovacion import global_route
    for t, skill, intent in [
        ("vigila la web https://ejemplo.com/pagina", "vigilancias", "web"),
        ("avísame si baja el precio de https://tienda.com/p", "vigilancias", "price"),
        ("avísame cuando haya noticias de zapatillas", "vigilancias", "news"),
        ("mis vigilancias", "vigilancias", "list"),
        ("borra la vigilancia 2", "vigilancias", "remove"),
        ("haz una copia de seguridad", "backup", "make"),
        ("qué copias de seguridad hay", "backup", "list"),
        ("archiva los bak", "backup", "baks"),
        ("mis informes", "research", "history"),
        ("abre el informe de calcetines", "research", "open_report"),
    ]:
        f, i = global_route(t)
        check((f, i) == (skill, intent), f"routing v19: '{t}' -> {f}/{i} (esperaba {skill}/{intent})")
    for t, skill in [
        ("qué me toca hoy", "coach"),
        ("recuérdame pagar al proveedor el viernes", "coach"),
        ("resume la pestaña 2", "chrome"),
        ("investiga el mercado de camisetas y hazme un informe", "research"),
        ("informe económico", "research"),
        ("qué tareas tengo", "tasks_board"),
        ("lee mis correos", "google_workspace"),
    ]:
        f, _ = global_route(t)
        check(f == skill, f"anti-robo v19: '{t}' -> {f} (esperaba {skill})")


if __name__ == "__main__":
    tests = [test_websearch_cache, test_websearch_extract, test_research_usa_motor_central,
             test_briefing_due, test_briefing_secciones,
             test_vigilancias_detectores, test_vigilancias_ciclo,
             test_informes_docx_e_historico, test_backup, test_routing_v19]
    for t in tests:
        print(f"· {t.__name__}")
        try:
            t()
        except Exception as e:
            _fail.append(f"{t.__name__} EXCEPCIÓN: {type(e).__name__}: {e}")
            print("  EXCEPCIÓN:", type(e).__name__, e)
    print(f"\n{'='*50}\n{_pass} pasados, {len(_fail)} fallados")
    sys.exit(1 if _fail else 0)
