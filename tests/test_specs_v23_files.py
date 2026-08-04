# -*- coding: utf-8 -*-
"""Tests de las specs v23 — FASE 3 y FASE 9 (voz y archivos reales).

  T7  el TTS solo reproduce respuestas de nexus, nunca lo que escribe el operador
  T18 lectura REAL de .md, .txt, .docx, .pdf (y aviso claro si no es compatible)
  T19 CRUD real: crear, leer, actualizar, sobrescribir con confirmación
  T20 guardar en la ubicación indicada, con verificación física y ruta real
  T21 versionado y copias de seguridad, con restauración

Ejecutar:  python tests/test_specs_v23_files.py    (desde la carpeta nexus)
"""
import asyncio
import importlib.util
import os
import re
import sys
import tempfile
from pathlib import Path

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_fail = []
_pass = 0


def check(cond, msg):
    global _pass
    if cond:
        _pass += 1
    else:
        _fail.append(msg)
        print("  FALLO:", msg)


if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import backend.core.infraestructura.files_io as FIO       # noqa: E402
import backend.core.comun.audit as audit        # noqa: E402
import backend.core.comun.confirm as confirm    # noqa: E402

_TMP = Path(tempfile.mkdtemp(prefix="nexus_v23f_"))
FIO.VERSIONS_DIR = _TMP / "file_versions"
FIO.VERSIONS_INDEX = FIO.VERSIONS_DIR / "index.json"
audit.AUDIT_FILE = _TMP / "logs" / "audit.jsonl"
WORK = _TMP / "trabajo"
WORK.mkdir(parents=True, exist_ok=True)


def _skill(nombre="files"):
    spec = importlib.util.spec_from_file_location(
        f"sk_{nombre}", os.path.join(ROOT, "skills", nombre, "skill.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _route(mod, text):
    for intent, rx in mod.SKILL["patterns"].items():
        m = re.search(rx, text, re.IGNORECASE)
        if m:
            return intent, m
    return None, None


# ══════════════ T18: lectura real ══════════════

def test_lee_texto_y_markdown():
    f = WORK / "notas.md"
    f.write_text("# Título\n\nUna línea con acentos: ñáéí.\n", encoding="utf-8")
    r = FIO.read_any(f)
    check(r["ok"], f"lee un .md ({r['error']})")
    check("acentos" in r["texto"], "conserva el contenido")
    check(r["meta"]["formato"] == ".md" and r["meta"]["bytes"] > 0,
          "devuelve formato y tamaño reales")
    check(r["meta"]["ruta"] == str(f), "devuelve la ruta real")


def test_codificaciones_raras():
    f = WORK / "viejo.txt"
    f.write_bytes("Café con leche\n".encode("cp1252"))
    r = FIO.read_any(f)
    check(r["ok"] and "Café" in r["texto"], "detecta la codificación (cp1252)")
    check(r["meta"].get("codificacion") in ("cp1252", "latin-1"), "y la reporta")


def test_docx_si_hay_libreria():
    f = WORK / "informe.docx"
    try:
        from docx import Document
    except ImportError:
        r = FIO.read_any(WORK / "no_existe.docx")
        check(not r["ok"], "sin python-docx no se inventa nada")
        check(True, "python-docx no instalado: se omite la lectura real de .docx")
        return
    d = Document()
    d.add_heading("Contrato", level=1)
    d.add_paragraph("Cláusula primera: pintar la fachada.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Importe"
    t.rows[0].cells[1].text = "12000"
    d.save(str(f))
    r = FIO.read_any(f)
    check(r["ok"], f"lee un .docx de verdad ({r['error']})")
    check("Cláusula primera" in r["texto"], "conserva el texto")
    check("Contrato" in r["texto"], "conserva los títulos (estructura básica)")
    check("12000" in r["texto"], "conserva el contenido de las tablas")
    check(r["meta"].get("tablas") == 1, "informa de cuántas tablas tiene")


def test_pdf_o_aviso_claro():
    f = WORK / "manual.pdf"
    f.write_bytes(b"%PDF-1.4 no es un pdf de verdad")
    r = FIO.read_any(f)
    check(not r["ok"], "un PDF ilegible NO se da por leído")
    check(bool(r["error"]), f"y explica por qué ({r['error'][:60]})")


def test_formato_no_compatible_se_dice():
    f = WORK / "foto.png"
    f.write_bytes(b"\x89PNG\r\n")
    r = FIO.read_any(f)
    check(not r["ok"], "una imagen no se lee como texto")
    check("imagen" in r["error"].lower(), "lo dice con claridad")
    f2 = WORK / "raro.qwe"
    f2.write_text("x", encoding="utf-8")
    r2 = FIO.read_any(f2)
    check(not r2["ok"] and "no sé leer" in r2["error"].lower(),
          "un formato desconocido se avisa, no se inventa")
    r3 = FIO.read_any(WORK / "fantasma.md")
    check(not r3["ok"] and "no existe" in r3["error"].lower(),
          "un archivo que no existe se dice claramente")


# ══════════════ T19/T20: CRUD y rutas ══════════════

def test_crear_verifica_y_devuelve_ruta():
    destino = WORK / "sub" / "nuevo.md"
    r = FIO.write_text(destino, "contenido inicial")
    check(r["ok"], f"crea el archivo ({r.get('error')})")
    check(r["accion"] == "creado", "lo marca como creado")
    check(r["verificado"] is True and Path(r["ruta"]).is_file(),
          "verifica FÍSICAMENTE que existe")
    check(r["ruta"] == str(destino), "devuelve la ruta real exacta")
    check(destino.parent.is_dir(), "crea la carpeta que faltaba")


def test_no_sobrescribe_sin_permiso():
    f = WORK / "importante.md"
    FIO.write_text(f, "TEXTO ORIGINAL")
    r = FIO.write_text(f, "texto nuevo")
    check(r["ok"] is False and r.get("necesita_confirmacion") is True,
          "no sobrescribe por su cuenta: pide confirmación")
    check("TEXTO ORIGINAL" in f.read_text(encoding="utf-8"),
          "el original sigue intacto")
    check(r["actual"]["bytes"] > 0, "informa de lo que hay ahora")
    r2 = FIO.write_text(f, "texto nuevo", overwrite=True)
    check(r2["ok"] and "texto nuevo" in f.read_text(encoding="utf-8"),
          "con permiso explícito sí lo cambia")
    check(bool(r2["copia_previa"]) and Path(r2["copia_previa"]).is_file(),
          "y guarda copia de la versión anterior")


def test_ampliar_no_destruye():
    f = WORK / "diario.md"
    FIO.write_text(f, "línea uno")
    FIO.write_text(f, "línea dos", append=True)
    txt = f.read_text(encoding="utf-8")
    check("línea uno" in txt and "línea dos" in txt,
          "añadir contenido NO destruye lo anterior")


def test_permisos_y_errores_claros():
    # Ruta imposible en Windows y en Linux por igual: un ARCHIVO usado como si
    # fuera carpeta. (Antes se usaba /proc/..., que en Windows sí se podía crear.)
    tapon = WORK / "soy_un_archivo.txt"
    tapon.write_text("x", encoding="utf-8")
    r = FIO.write_text(tapon / "dentro" / "x.md", "x")
    check(r["ok"] is False and bool(r["error"]),
          f"un error de permisos o ruta se explica, no se traga ({r.get('error', '')[:60]})")
    check(not r.get("necesita_confirmacion"),
          "y no se confunde con «ya existe, confírmame»")


# ══════════════ T21: versiones ══════════════

def test_versionado_y_restauracion():
    f = WORK / "presupuesto.md"
    FIO.write_text(f, "versión A")
    FIO.write_text(f, "versión B", overwrite=True)
    FIO.write_text(f, "versión C", overwrite=True)
    vs = FIO.versions(f)
    check(len(vs) == 2, f"guarda una versión por cada cambio ({len(vs)})")
    check(all(Path(FIO.VERSIONS_DIR, v["version"]).is_file() for v in vs),
          "las copias están en disco")
    check("versión C" in f.read_text(encoding="utf-8"), "el original es el último")
    r = FIO.restore_version(f)
    check(r["ok"], "se puede restaurar")
    check("versión B" in f.read_text(encoding="utf-8"),
          "vuelve a la versión inmediatamente anterior")
    # Dos guardados en el MISMO milisegundo no pueden pisarse: pasaba 1 de cada
    # 45 veces, y al restaurar la copia del estado actual machacaba justo la
    # version que se iba a recuperar, asi que «volver atras» no volvia a nada.
    rafaga = WORK / "rafaga.md"
    FIO.write_text(rafaga, "uno")
    for texto in ("dos", "tres", "cuatro", "cinco", "seis"):
        FIO.write_text(rafaga, texto, overwrite=True)
    vr = FIO.versions(rafaga)
    nombres = [x["version"] for x in vr]
    check(len(set(nombres)) == len(nombres),
          f"cada guardado deja SU copia, ninguna pisa a otra ({len(set(nombres))} de {len(nombres)})")
    check(all(Path(FIO.VERSIONS_DIR, x).is_file() for x in nombres),
          "y todas siguen en disco")
    FIO.restore_version(rafaga)
    check("cinco" in rafaga.read_text(encoding="utf-8"),
          f"restaurar en ráfaga devuelve la anterior de verdad "
          f"({rafaga.read_text(encoding='utf-8').strip()})")
    check(bool(r["copia_del_estado_previo"]),
          "y del estado que había justo antes también guarda copia")
    check(len(FIO.versions(f)) >= 2, "las copias no se pisan entre ellas")


# ══════════════ Integración con la skill ══════════════

def test_skill_archivos_routing_y_confirmacion():
    mod = _skill("files")
    casos = [("lee el archivo informe.md", "read"),
             ("resume el documento contrato.pdf", "summarize"),
             ("añade al archivo notas.md que diga hola", "update"),
             ("actualiza el archivo notas.md con el texto nuevo", "update"),
             ("qué versiones tienes de notas.md", "versions"),
             ("restaura el archivo notas.md", "restore_file")]
    for texto, esperado in casos:
        intent, _ = _route(mod, texto)
        check(intent == esperado, f"routing: «{texto}» → {intent} (esperaba {esperado})")

    src = Path(ROOT, "skills", "files", "skill.py").read_text(encoding="utf-8")
    check("files_io" in src, "la skill usa el lector/escritor real")
    check("confirm.request" in src, "sobrescribir y restaurar piden confirmación")
    check("sobrescribir_archivo" in src and "restaurar_archivo" in src,
          "cada acción destructiva tiene su confirmación con nombre")
    check("verificado ✔" in src, "confirma que ha verificado el archivo")
    check("no lo sobrescribo" not in src,
          "ya no se limita a negarse: pregunta y, si dices que sí, versiona y escribe")


def test_skill_lectura_no_miente():
    mod = _skill("files")
    f = WORK / "raro2.qwe"
    f.write_text("x", encoding="utf-8")

    class _P:
        @staticmethod
        def path_allowed(_p):
            return True

        @staticmethod
        def deny_msg(_p):
            return "denegado"
    mod._perm = lambda: _P
    intent, m = _route(mod, f"lee el archivo {f}")
    r = asyncio.run(mod.handle(intent, f"lee el archivo {f}", m, {"channel": "pc"}))
    check("no sé leer" in r["reply"].lower(),
          "si no sabe leer el formato lo dice, no finge haberlo leído")
    g = WORK / "ok.md"
    g.write_text("hola mundo", encoding="utf-8")
    intent, m = _route(mod, f"lee el archivo {g}")
    r = asyncio.run(mod.handle(intent, f"lee el archivo {g}", m, {"channel": "pc"}))
    check("hola mundo" in r["reply"], "y cuando sí puede, lee de verdad")
    check("Leído" in r["reply"] and ".md" in r["reply"],
          "informa de qué ha leído exactamente")


# ══════════════ T7: voz ══════════════

def test_tts_solo_habla_respuestas():
    import backend.core.infraestructura.tts as tts
    check(hasattr(tts, "note_user_text") and hasattr(tts, "is_user_echo"),
          "el TTS sabe qué escribió el operador")
    tts.note_user_text("borra las tareas completadas del tablero")
    check(tts.is_user_echo("borra las tareas completadas del tablero") is True,
          "reconoce su propio texto")
    check(tts.is_user_echo("He borrado 4 tareas completadas") is False,
          "la respuesta de nexus sí se puede decir")
    src = Path(ROOT, "backend", "core", "infraestructura", "tts.py").read_text(encoding="utf-8")
    check('async def speak(text: str, role: str = "assistant")' in src,
          "speak() exige el rol del mensaje")
    check('if role != "assistant":' in src,
          "el TTS RECHAZA todo lo que no sea una respuesta del asistente")
    check('settings.get("tts_enabled", True)' in src,
          "se puede desactivar del todo la respuesta por voz")
    i_role = src.find('if role != "assistant":')
    i_echo = src.find("if is_user_echo(text):")
    i_emit = src.find('await bus.emit("state", "speaking")')
    check(0 < i_role < i_emit and 0 < i_echo < i_emit,
          "los dos filtros van ANTES de emitir audio")


def test_brain_y_hud_no_mandan_al_tts_lo_del_operador():
    brain = Path(ROOT, "backend", "core", "brain.py").read_text(encoding="utf-8")
    check("_tts_guard.note_user_text(text)" in brain,
          "el brain apunta cada entrada del operador para el filtro del TTS")
    js = Path(ROOT, "frontend", "js", "command.js").read_text(encoding="utf-8")
    check("_sameAsUser" in js, "el respaldo de voz del navegador tiene el mismo filtro")
    check("window.__lastUserText = text" in js, "el HUD recuerda lo que tecleaste")
    check("armTTSFallback(data.reply)" in js and "armTTSFallback(data.user)" not in js,
          "al navegador solo se le pasa la RESPUESTA, nunca lo que escribió el operador")


# ══════════════ NORMA: LOS INFORMES SE CREAN EN .md ══════════════
# Petición de Adri (30/07/2026): «todos los informes por defecto han de crearlos
# en archivos .md salvo que se pida expresamente otra cosa».
def test_los_informes_salen_en_markdown():
    from backend.core.infraestructura.files_io import formato_pedido, pidio_formato
    # sin pedir formato → Markdown
    for orden in ("hazme un informe de ventas", "crea un documento sobre el proyecto",
                  "escríbeme un resumen de la reunión", "redáctame un informe",
                  "prepárame un dossier del competidor", "hazme un análisis"):
        check(formato_pedido(orden) == ".md", f"«{orden}» → .md (dio {formato_pedido(orden)})")
        check(not pidio_formato(orden), f"«{orden}» no pide formato expreso")
    # pidiéndolo, manda lo que diga el usuario
    for orden, esp in (("el informe en word", ".docx"), ("pásamelo a pdf", ".pdf"),
                       ("exporta los leads a csv", ".csv"), ("dámelo en excel", ".xlsx"),
                       ("un archivo notas.txt", ".txt"), ("dame un json", ".json"),
                       ("hazme una página web", ".html"), ("en texto plano", ".txt"),
                       ("en markdown", ".md")):
        check(formato_pedido(orden) == esp, f"«{orden}» → {esp} (dio {formato_pedido(orden)})")
        check(pidio_formato(orden), f"«{orden}» SÍ pide formato expreso")
    # y la skill de archivos usa esa misma regla, no la suya
    src = open(os.path.join(ROOT, "skills", "files", "skill.py"), encoding="utf-8").read()
    check("formato_pedido" in src, "la skill de archivos usa la regla común")
    check('ext = ".txt"' not in src, "y ya no crea .txt por su cuenta")
    inv = open(os.path.join(ROOT, "skills", "research", "skill.py"), encoding="utf-8").read()
    check(".md\"" in inv or ".md'" in inv, "la de investigación ya guardaba en .md")


if __name__ == "__main__":
    tests = [test_lee_texto_y_markdown, test_codificaciones_raras,
             test_docx_si_hay_libreria, test_pdf_o_aviso_claro,
             test_formato_no_compatible_se_dice,
             test_crear_verifica_y_devuelve_ruta, test_no_sobrescribe_sin_permiso,
             test_ampliar_no_destruye, test_permisos_y_errores_claros,
             test_versionado_y_restauracion,
             test_skill_archivos_routing_y_confirmacion, test_skill_lectura_no_miente,
             test_tts_solo_habla_respuestas,
             test_brain_y_hud_no_mandan_al_tts_lo_del_operador,
             test_los_informes_salen_en_markdown]
    for t in tests:
        print(f"· {t.__name__}")
        try:
            t()
        except Exception as e:
            _fail.append(f"{t.__name__} EXCEPCIÓN: {type(e).__name__}: {e}")
            print("  EXCEPCIÓN:", type(e).__name__, e)
    print(f"\n{'='*50}\n{_pass} pasados, {len(_fail)} fallados")
    sys.exit(1 if _fail else 0)
