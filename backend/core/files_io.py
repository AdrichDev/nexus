"""
nexus — LECTURA Y ESCRITURA REAL DE ARCHIVOS (specs v23, TAREAS 18-21).

Reglas fijadas con Adri:
  * Se leen archivos DE VERDAD: .md, .txt, .docx y .pdf como mínimo (y .json,
    .csv, .xlsx y código si hay librería). Si un formato no es compatible, se
    DICE — nunca se afirma haber leído algo que no se ha abierto.
  * Antes de modificar o sobrescribir un archivo existente se guarda una VERSIÓN
    anterior en data/file_versions/, y se puede restaurar.
  * Después de crear o modificar se VERIFICA físicamente que el archivo existe y
    se devuelve su RUTA REAL. Nada de rutas inventadas.
  * El borrado definitivo no vive aquí: va por la papelera del sistema y con
    confirmación (backend/core/confirm.py).
"""
from __future__ import annotations

import datetime as dt
import re
import json
import os
import shutil
from pathlib import Path

from .config import DATA_DIR, CONFIG_DIR

VERSIONS_DIR = DATA_DIR / "file_versions"
VERSIONS_INDEX = VERSIONS_DIR / "index.json"

TEXTO = {".txt", ".md", ".markdown", ".json", ".csv", ".tsv", ".log", ".ini",
         ".cfg", ".yaml", ".yml", ".xml", ".html", ".htm", ".py", ".js", ".ts",
         ".jsx", ".tsx", ".java", ".go", ".rs", ".c", ".cpp", ".h", ".cs",
         ".php", ".rb", ".sql", ".sh", ".ps1", ".bat", ".env"}
BINARIOS_CONOCIDOS = {".docx", ".pdf", ".xlsx", ".pptx"}
IMAGENES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg"}

_MAX_CHARS = 200_000


def _now_tag() -> str:
    """Con milisegundos: dos copias en el mismo segundo NO pueden pisarse
    (lo cazó el test de versionado — la segunda machacaba a la primera)."""
    return dt.datetime.now().strftime("%Y%m%d-%H%M%S-%f")[:-3]


def _audit(**kw) -> None:
    try:
        from . import audit as _a
        _a.log(**kw)
    except Exception:
        pass


# ───────────────────────── LECTURA (T18) ─────────────────────────

def _leer_docx(path: Path) -> tuple[str, dict]:
    try:
        from docx import Document                      # python-docx
    except ImportError:
        raise RuntimeError("para leer .docx falta la librería python-docx "
                           "(pip install python-docx)")
    doc = Document(str(path))
    partes, tablas = [], 0
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        est = (p.style.name or "") if p.style else ""
        partes.append(("#" * min(4, int(est[-1])) + " " + t) if est.startswith("Heading")
                      and est[-1].isdigit() else t)
    for tb in doc.tables:
        tablas += 1
        for fila in tb.rows:
            partes.append(" | ".join((c.text or "").strip() for c in fila.cells))
    return "\n".join(partes), {"parrafos": len(doc.paragraphs), "tablas": tablas}


def _leer_pdf(path: Path) -> tuple[str, dict]:
    texto, paginas = "", 0
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            paginas = len(pdf.pages)
            texto = "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except ImportError:
        try:
            from pypdf import PdfReader
        except ImportError:
            try:
                from PyPDF2 import PdfReader          # type: ignore
            except ImportError:
                raise RuntimeError("para leer PDF falta una librería de PDF "
                                   "(pip install pdfplumber  ·  o pypdf)")
        r = PdfReader(str(path))
        paginas = len(r.pages)
        texto = "\n\n".join((pg.extract_text() or "") for pg in r.pages)
    if not texto.strip():
        raise RuntimeError("el PDF no tiene texto extraíble (parece escaneado): "
                           "haría falta OCR, y eso todavía no lo hago")
    return texto, {"paginas": paginas}


def _umbral_xlsx_aviso() -> int:
    """Filas por encima de las cuales _leer_xlsx() avisa (pero NO corta).
    002-memoria-y-conocimiento (bloque C, C1.4): antes de esto, _leer_xlsx()
    cortaba de un corte mudo a 3000 filas por hoja -- un .xlsx real (p.ej.
    «Wabiks Content Intelligence.xlsx») podía perder filas sin que nadie se
    enterase. Ahora se lee entero siempre; esto solo decide cuándo avisar."""
    try:
        f = CONFIG_DIR / "umbrales.json"
        if f.is_file():
            v = (json.loads(f.read_text(encoding="utf-8")) or {}).get(
                "memoria", {}).get("ingesta", {}).get("aviso_filas_xlsx")
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                return int(v)
    except Exception:
        pass
    return 50_000


def _leer_xlsx(path: Path) -> tuple[str, dict]:
    """Lee el .xlsx ENTERO, sin cortar filas. Antes (bloque C, C1.4) cortaba
    a 3000 filas por hoja con un «… (hoja truncada)» que nadie leía: un
    catálogo real de más de 3000 filas entraba mutilado en la memoria."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("para leer .xlsx falta openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        partes, filas = [], 0
        for hoja in wb.worksheets:
            partes.append(f"## Hoja: {hoja.title}")
            for fila in hoja.iter_rows(values_only=True):
                filas += 1
                partes.append(" | ".join("" if c is None else str(c) for c in fila))
        aviso = _umbral_xlsx_aviso()
        if filas > aviso:
            partes.append(f"… (aviso: {filas} filas leídas — por encima del umbral de aviso "
                           f"{aviso}, pero NADA se ha cortado)")
        n_hojas = len(wb.worksheets)
    finally:
        # INCIDENTE (bloque C): en modo read_only, openpyxl deja el fichero
        # ABIERTO hasta wb.close() -- sin esto, el buzón no podía mover el
        # .xlsx ya leído a ingested/ (WinError 32, en uso por otro proceso).
        wb.close()
    return "\n".join(partes), {"hojas": n_hojas, "filas": filas}


def leer_xlsx_estructurado(path: Path) -> list[dict]:
    """Lee un .xlsx preservando su ESTRUCTURA hoja/columna/fila, sin
    aplanarlo a texto. Devuelve [{"hoja", "cabeceras", "filas"}] por hoja,
    con `filas` como lista de tuplas alineadas con `cabeceras`.
    002-memoria-y-conocimiento: memoria-ingesta-documentos, requisito
    «.xlsx por hojas y columnas» — trocear por caracteres partía filas por
    la mitad y perdía a qué columna pertenecía cada valor."""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("para leer .xlsx falta openpyxl (pip install openpyxl)")
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        hojas = []
        for hoja in wb.worksheets:
            filas_it = hoja.iter_rows(values_only=True)
            cabeceras_raw = next(filas_it, None) or ()
            cabeceras = [str(c) if c is not None else f"columna_{i + 1}"
                         for i, c in enumerate(cabeceras_raw)]
            filas = [fila for fila in filas_it if any(v is not None for v in fila)]
            hojas.append({"hoja": hoja.title, "cabeceras": cabeceras, "filas": filas})
    finally:
        wb.close()  # ver comentario del mismo incidente en _leer_xlsx()
    return hojas


def puede_leer(path: Path) -> tuple[bool, str]:
    """¿Sé abrir esto? (True, '') o (False, motivo claro)."""
    ext = path.suffix.lower()
    if ext in TEXTO or ext == "":
        return True, ""
    if ext in BINARIOS_CONOCIDOS:
        return True, ""
    if ext in IMAGENES:
        return False, (f"«{path.name}» es una imagen: puedo decirte que existe y cuánto "
                       "ocupa, pero todavía no leo su contenido.")
    return False, (f"No sé leer el formato «{ext or 'sin extensión'}» de «{path.name}». "
                   "Formatos que sí leo: .md, .txt, .docx, .pdf, .json, .csv, .xlsx y código.")


def read_any(path, *, limite: int | None = None) -> dict:
    """Lee un archivo REAL. Devuelve
    {ok, texto, meta:{formato, bytes, modificado, truncado, …}, error}.
    Nunca lanza: siempre hay algo que contarle al operador.

    `limite` (002-memoria-y-conocimiento, bloque C, C1.3):
    - None (por defecto): comportamiento IDÉNTICO al de siempre, corta a
      _MAX_CHARS. Así ningún llamador existente (skills/files/skill.py,
      tests/test_specs_v23_files.py) cambia de comportamiento.
    - 0: sin límite -- el texto entra ENTERO. Lo usa la ingesta de
      documentos, que necesita el documento completo o nada (rechaza si
      meta['truncado'] sale True).
    - N > 0: límite propio en caracteres.
    """
    p = Path(path).expanduser()
    if not p.exists():
        return {"ok": False, "texto": "", "meta": {},
                "error": f"No existe la ruta {p}."}
    if p.is_dir():
        return {"ok": False, "texto": "", "meta": {},
                "error": f"{p} es una carpeta, no un archivo."}
    ok, motivo = puede_leer(p)
    if not ok:
        return {"ok": False, "texto": "", "meta": {"formato": p.suffix.lower()},
                "error": motivo}
    ext = p.suffix.lower()
    meta = {"formato": ext or "texto", "bytes": p.stat().st_size, "ruta": str(p),
            "modificado": dt.datetime.fromtimestamp(p.stat().st_mtime).isoformat(
                timespec="seconds")}
    try:
        if ext == ".docx":
            texto, extra = _leer_docx(p)
        elif ext == ".pdf":
            texto, extra = _leer_pdf(p)
        elif ext == ".xlsx":
            texto, extra = _leer_xlsx(p)
        else:
            data = p.read_bytes()
            for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
                try:
                    texto = data.decode(enc)
                    extra = {"codificacion": enc}
                    break
                except UnicodeDecodeError:
                    continue
            else:
                return {"ok": False, "texto": "", "meta": meta,
                        "error": f"No puedo descifrar la codificación de {p.name}."}
    except Exception as exc:                            # noqa: BLE001
        return {"ok": False, "texto": "", "meta": meta, "error": str(exc)}
    meta.update(extra)
    meta["caracteres"] = len(texto)
    if limite is None:
        tope = _MAX_CHARS
    elif limite == 0:
        tope = None
    else:
        tope = int(limite)
    if tope is not None and len(texto) > tope:
        texto_final = texto[:tope]
        meta["truncado"] = True
    else:
        texto_final = texto
        meta["truncado"] = False
    _audit(action="file_read", destructive=False, request=str(p),
           result=f"{meta['formato']} · {meta['caracteres']} caracteres"
                  f"{' · TRUNCADO' if meta['truncado'] else ''}")
    return {"ok": True, "texto": texto_final, "meta": meta, "error": ""}


# ───────────────────────── VERSIONES (T21) ─────────────────────────

def _index_load() -> dict:
    try:
        return json.loads(VERSIONS_INDEX.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _index_save(d: dict) -> None:
    VERSIONS_DIR.mkdir(parents=True, exist_ok=True)
    VERSIONS_INDEX.write_text(json.dumps(d, ensure_ascii=False, indent=1),
                              encoding="utf-8")


def _copia_libre(destino: Path) -> Path:
    """Un nombre de copia que NO pise a otra copia.

    El sello de tiempo llega al milisegundo, y dos guardados seguidos caen en el
    mismo milisegundo mas a menudo de lo que parece (1 de cada 45 en una prueba
    de 400). Cuando pasaba, la copia nueva SOBRESCRIBIA a la anterior: se perdia
    una version en silencio y —peor— al restaurar, la copia de seguridad del
    estado actual machacaba justo la version que se iba a recuperar, asi que
    «volver atras» devolvia el archivo tal y como estaba."""
    if not destino.exists():
        return destino
    for n in range(1, 1000):
        alterna = destino.with_name(f"{destino.stem}-{n}{destino.suffix}")
        if not alterna.exists():
            return alterna
    return destino


def backup(path) -> str:
    """Guarda una COPIA de la versión actual antes de tocarla. Devuelve la ruta
    de la copia ('' si el archivo no existía todavía)."""
    p = Path(path).expanduser()
    if not p.is_file():
        return ""
    VERSIONS_DIR.mkdir(parents=True, exist_ok=True)
    destino = _copia_libre(VERSIONS_DIR / f"{p.stem}.{_now_tag()}{p.suffix}.bak")
    shutil.copy2(str(p), str(destino))
    idx = _index_load()
    idx.setdefault(str(p), []).append(
        {"version": destino.name, "guardada": dt.datetime.now().isoformat(timespec="seconds"),
         "bytes": destino.stat().st_size})
    idx[str(p)] = idx[str(p)][-20:]                 # 20 versiones por archivo
    _index_save(idx)
    _audit(action="file_backup", destructive=False, request=str(p),
           result=f"versión anterior guardada en {destino}")
    return str(destino)


def versions(path) -> list[dict]:
    return list(reversed(_index_load().get(str(Path(path).expanduser()), [])))


def restore_version(path, version: str = "") -> dict:
    """Devuelve el archivo a una versión anterior (la última si no se dice cuál).
    La copia NUNCA sobrescribe sin dejar antes copia del estado actual."""
    p = Path(path).expanduser()
    vs = versions(p)
    if not vs:
        return {"ok": False, "error": f"No tengo versiones guardadas de {p.name}."}
    elegida = next((v for v in vs if v["version"] == version), vs[0]) if version else vs[0]
    origen = VERSIONS_DIR / elegida["version"]
    if not origen.is_file():
        return {"ok": False, "error": f"La copia {elegida['version']} ya no está en disco."}
    previa = backup(p)
    shutil.copy2(str(origen), str(p))
    _audit(action="file_restore", destructive=True, confirmed=True, request=str(p),
           result=f"restaurada la versión {elegida['version']}")
    return {"ok": True, "ruta": str(p), "version": elegida["version"],
            "copia_del_estado_previo": previa,
            "guardada": elegida.get("guardada", "")}


# ───────────────────────── ESCRITURA (T19/T20) ─────────────────────────

def verify(path) -> dict:
    p = Path(path).expanduser()
    if not p.exists():
        return {"existe": False, "ruta": str(p)}
    st = p.stat()
    return {"existe": True, "ruta": str(p), "bytes": st.st_size,
            "modificado": dt.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")}


def write_text(path, contenido: str, *, overwrite: bool = False,
               append: bool = False, reason: str = "") -> dict:
    """Crea o modifica un archivo de texto. SIEMPRE:
      * crea la carpeta si hace falta,
      * hace copia de la versión anterior si el archivo existía,
      * verifica FÍSICAMENTE que quedó escrito,
      * devuelve la RUTA REAL.
    Si el archivo existe y no se pide overwrite/append, NO lo toca: devuelve
    {'ok': False, 'necesita_confirmacion': True} para que arriba se pregunte."""
    p = Path(path).expanduser()
    existia = p.is_file()
    if existia and not (overwrite or append):
        return {"ok": False, "necesita_confirmacion": True, "ruta": str(p),
                "error": f"«{p.name}» ya existe en {p.parent}.",
                "actual": verify(p)}
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
    except Exception as exc:                            # noqa: BLE001
        return {"ok": False, "ruta": str(p),
                "error": f"No puedo crear la carpeta {p.parent}: {exc}"}
    if not os.access(str(p.parent), os.W_OK):
        return {"ok": False, "ruta": str(p),
                "error": f"No tengo permiso de escritura en {p.parent}."}
    copia = backup(p) if existia else ""
    try:
        if append and existia:
            previo = p.read_text(encoding="utf-8", errors="replace")
            nuevo = previo + ("" if previo.endswith("\n") else "\n") + contenido
        else:
            nuevo = contenido
        p.write_text(nuevo if nuevo.endswith("\n") else nuevo + "\n", encoding="utf-8")
    except Exception as exc:                            # noqa: BLE001
        return {"ok": False, "ruta": str(p), "copia_previa": copia,
                "error": f"No he podido escribir {p}: {exc}"}
    info = verify(p)                                    # verificación FÍSICA
    if not info["existe"]:
        return {"ok": False, "ruta": str(p),
                "error": "Escribí sin error pero el archivo no aparece en disco."}
    accion = "modificado" if existia else "creado"
    _audit(action=f"file_{accion}", destructive=existia, confirmed=True,
           request=reason or str(p), result=f"{accion}: {p} ({info['bytes']} bytes)",
           extra={"copia_previa": copia})
    return {"ok": True, "ruta": str(p), "accion": accion, "copia_previa": copia,
            "bytes": info["bytes"], "verificado": True}


def resolve_target(nombre: str, base) -> Path:
    """Ruta final a partir de un nombre y una carpeta base ya resuelta."""
    n = (nombre or "").strip().strip('"\'')
    p = Path(n).expanduser()
    return p if p.is_absolute() else Path(base).expanduser() / n

# ── QUÉ FORMATO USAR CUANDO NADIE LO DICE ────────────────────────────────────
# Norma de Adri (30/07/2026): «todos los informes por defecto han de crearse en
# archivos .md salvo que se pida expresamente otra cosa».
#
# Está AQUÍ, en un solo sitio, a propósito: antes cada skill decidía por su
# cuenta (la de archivos escribía .txt, la de investigación .md) y bastaba con
# que naciera una skill nueva para volver a tener tres criterios distintos.
# Quien tenga que crear un documento pregunta a esta función y se acabó.
FORMATOS = {
    "md": (r"\bmarkdown\b|\.md\b|\bformato\s+md\b", ".md"),
    "docx": (r"\bword\b|\.docx\b|\bdocumento\s+de\s+word\b", ".docx"),
    "pdf": (r"\bpdf\b|\.pdf\b", ".pdf"),
    "xlsx": (r"\bexcel\b|\.xlsx\b|\bhoja\s+de\s+c[aá]lculo\b", ".xlsx"),
    "csv": (r"\bcsv\b|\.csv\b", ".csv"),
    "txt": (r"\btexto\s+plano\b|\.txt\b|\bbloc\s+de\s+notas\b|\ben\s+txt\b", ".txt"),
    "json": (r"\bjson\b|\.json\b", ".json"),
    "html": (r"\bhtml\b|\.html?\b|\bp[aá]gina\s+web\b", ".html"),
}
_FORMATOS_RX = [(re.compile(pat, re.IGNORECASE), ext) for pat, ext in FORMATOS.values()]


def formato_pedido(texto: str, defecto: str = ".md") -> str:
    """La extensión que ha pedido el usuario, o «.md» si no ha pedido ninguna.

    >>> formato_pedido("hazme un informe de ventas")
    '.md'
    >>> formato_pedido("hazme el informe en word")
    '.docx'
    """
    t = texto or ""
    for rx, ext in _FORMATOS_RX:
        if rx.search(t):
            return ext
    return defecto


def pidio_formato(texto: str) -> bool:
    """¿Ha dicho el usuario EXPRESAMENTE en qué formato lo quiere?"""
    return formato_pedido(texto, defecto="") != ""
